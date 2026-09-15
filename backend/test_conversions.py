import json
import os
import shutil
import tempfile
from pathlib import Path
from PIL import Image
import docx
import pandas as pd
import piexif

from backend.executor import execute_job
from backend.app.services.scanner import scan_text_for_pii, sanitize_text
from backend.app.services.signer import create_download_token, verify_download_token
from backend.app.services.mime import sniff_mime

def run_all_tests():
    test_dir = Path(tempfile.mkdtemp(prefix="fluxdrive_test_"))
    print(f"Running verification tests in temporary sandbox: {test_dir}")
    
    passed = 0
    total = 0

    def check(name: str, cond: bool, extra=""):
        nonlocal passed, total
        total += 1
        if cond:
            passed += 1
            print(f"  [PASS] {name} {extra}")
        else:
            print(f"  [FAIL] {name} {extra}")

    try:
        # 1. Create Base Test Files
        # Sample PNG
        png_path = test_dir / "sample.png"
        img = Image.new("RGBA", (100, 100), color=(255, 0, 0, 200))
        img.save(png_path, "PNG")

        # Sample JPG
        jpg_path = test_dir / "sample.jpg"
        img_rgb = Image.new("RGB", (100, 100), color=(0, 255, 0))
        img_rgb.save(jpg_path, "JPEG")

        # Sample WEBP
        webp_path = test_dir / "sample.webp"
        img_rgb.save(webp_path, "WEBP")

        # Sample CSV
        csv_path = test_dir / "sample.csv"
        csv_path.write_text("id,name,score\n1,Alice,95\n2,Bob,88\n", encoding="utf-8")

        # Sample JSON
        json_path = test_dir / "sample.json"
        json_path.write_text(json.dumps([{"id": 1, "product": "Widget", "price": 25.5}, {"id": 2, "product": "Gear", "price": 10.0}]), encoding="utf-8")

        # Sample DOCX
        docx_path = test_dir / "sample.docx"
        doc = docx.Document()
        doc.add_heading("FluxDrive Security Architecture", level=1)
        doc.add_paragraph("This document outlines the privacy-first ephemeral file workspace.")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Metric"
        t.cell(0, 1).text = "Status"
        t.cell(1, 0).text = "PII Protection"
        t.cell(1, 1).text = "Active"
        doc.save(docx_path)

        out_dir = test_dir / "outputs"
        out_dir.mkdir()

        # Test Conversions
        # 1. PNG -> JPG
        res = execute_job({"operation": "png_to_jpg", "input_paths": [str(png_path)], "output_dir": str(out_dir)})
        check("PNG -> JPG", res["success"] and Path(res["output_path"]).exists())

        # 2. PNG -> WEBP
        res = execute_job({"operation": "png_to_webp", "input_paths": [str(png_path)], "output_dir": str(out_dir)})
        check("PNG -> WEBP", res["success"] and Path(res["output_path"]).exists())

        # 3. JPG -> PNG
        res = execute_job({"operation": "jpg_to_png", "input_paths": [str(jpg_path)], "output_dir": str(out_dir)})
        check("JPG -> PNG", res["success"] and Path(res["output_path"]).exists())

        # 4. JPG -> WEBP
        res = execute_job({"operation": "jpg_to_webp", "input_paths": [str(jpg_path)], "output_dir": str(out_dir)})
        check("JPG -> WEBP", res["success"] and Path(res["output_path"]).exists())

        # 5. WEBP -> PNG
        res = execute_job({"operation": "webp_to_png", "input_paths": [str(webp_path)], "output_dir": str(out_dir)})
        check("WEBP -> PNG", res["success"] and Path(res["output_path"]).exists())

        # 6. WEBP -> JPG
        res = execute_job({"operation": "webp_to_jpg", "input_paths": [str(webp_path)], "output_dir": str(out_dir)})
        check("WEBP -> JPG", res["success"] and Path(res["output_path"]).exists())

        # 7. Image -> PDF
        res = execute_job({"operation": "image_to_pdf", "input_paths": [str(png_path)], "output_dir": str(out_dir)})
        pdf_out = Path(res["output_path"]) if res["success"] else None
        check("PNG -> PDF", res["success"] and pdf_out and pdf_out.exists())

        # 8. Batch Images -> Multi-page PDF
        res = execute_job({"operation": "batch_to_pdf", "input_paths": [str(png_path), str(jpg_path), str(webp_path)], "output_dir": str(out_dir)})
        batch_pdf = Path(res["output_path"]) if res["success"] else None
        check("Batch Images -> Multi-page PDF", res["success"] and batch_pdf and batch_pdf.exists())

        # 9. PDF -> PNG
        if pdf_out:
            res = execute_job({"operation": "pdf_to_png", "input_paths": [str(pdf_out)], "output_dir": str(out_dir)})
            check("PDF -> PNG", res["success"] and Path(res["output_path"]).exists())

        # 10. PDF -> JPG
        if pdf_out:
            res = execute_job({"operation": "pdf_to_jpg", "input_paths": [str(pdf_out)], "output_dir": str(out_dir)})
            check("PDF -> JPG", res["success"] and Path(res["output_path"]).exists())

        # 11. CSV -> JSON
        res = execute_job({"operation": "csv_to_json", "input_paths": [str(csv_path)], "output_dir": str(out_dir)})
        check("CSV -> JSON", res["success"] and Path(res["output_path"]).exists())

        # 12. JSON -> CSV
        res = execute_job({"operation": "json_to_csv", "input_paths": [str(json_path)], "output_dir": str(out_dir)})
        check("JSON -> CSV", res["success"] and Path(res["output_path"]).exists())

        # 13. DOCX -> PDF
        res = execute_job({"operation": "docx_to_pdf", "input_paths": [str(docx_path)], "output_dir": str(out_dir)})
        check("DOCX -> PDF", res["success"] and Path(res["output_path"]).exists())

        # 14. Compress Image (real disk byte reduction)
        # Create a large uncompressed test image
        large_img_path = test_dir / "large_raw.png"
        large_img = Image.new("RGB", (800, 800), color=(120, 180, 240))
        for x in range(0, 800, 20):
            for y in range(0, 800, 20):
                large_img.putpixel((x, y), (x % 256, y % 256, (x+y) % 256))
        large_img.save(large_img_path, "PNG")
        
        res = execute_job({"operation": "compress_image", "input_paths": [str(large_img_path)], "output_dir": str(out_dir), "options": {"quality": 40}})
        savings = res.get("metrics", {}).get("savings_pct", 0)
        check("Compress Image (disk savings)", res["success"] and savings > 0, f"(savings: {savings}%)")

        # 15. Strip Metadata
        # Create image with real EXIF
        exif_img_path = test_dir / "with_exif.jpg"
        im_exif = Image.new("RGB", (200, 200), color=(100, 100, 100))
        zeroth_ifd = {
            piexif.ImageIFD.Make: "FluxTestCamera",
            piexif.ImageIFD.Model: "FluxPro-X1",
            piexif.ImageIFD.Software: "TestFirmware 1.0"
        }
        exif_dict = {"0th": zeroth_ifd, "Exif": {}, "GPS": {}, "1st": {}, "thumbnail": None}
        exif_bytes = piexif.dump(exif_dict)
        im_exif.save(exif_img_path, "JPEG", exif=exif_bytes)

        res = execute_job({"operation": "strip_metadata", "input_paths": [str(exif_img_path)], "output_dir": str(out_dir)})
        stripped_file = Path(res["output_path"]) if res["success"] else None
        # Verify stripped file has NO EXIF
        has_exif_after = False
        if stripped_file:
            with Image.open(stripped_file) as check_img:
                raw_ex = check_img.getexif()
                has_exif_after = bool(raw_ex)
        check("Strip Metadata (clean output verified)", res["success"] and not has_exif_after)

        # 16. Privacy Scanner Test
        synthetic_pii_text = """
        Confidential User Report:
        Client: Jane Doe
        Email: jane.doe@acmeworks.org
        Contact: 415-555-2671
        Aadhaar ID: 5412 8901 2345
        PAN Number: ABCDE1234F
        Payment Card: 4532 0150 0000 0007
        Local Server: 192.168.1.105
        Portal: https://internal.acme.corp/dashboard
        """
        scan_result = scan_text_for_pii(synthetic_pii_text)
        types_found = {f["type"] for f in scan_result["findings"]}
        expected_types = {"email", "phone", "aadhaar", "pan", "credit_card", "ipv4", "url"}
        all_pii_found = expected_types.issubset(types_found)
        check("Privacy Scanner (All 7 PII categories + Luhn verify)", all_pii_found, f"Found: {types_found}")

        # 17. Sanitization Masking Test
        sanitized_txt, masked_count = sanitize_text(synthetic_pii_text)
        check("Sanitization Masking", masked_count >= 5 and "jane.doe@acmeworks.org" not in sanitized_txt)

        # 18. Signed Expiring Token Test
        token = create_download_token("file123", "test.pdf", expires_in_seconds=300)
        verified = verify_download_token(token)
        check("HMAC Signed Download Token (Valid)", verified is not None and verified["file_id"] == "file123")

        # Expired token rejection
        expired_token = create_download_token("file123", "test.pdf", expires_in_seconds=-10)
        verified_expired = verify_download_token(expired_token)
        check("HMAC Signed Download Token (Rejects Expired)", verified_expired is None)

        # Tampered token rejection
        tampered_token = token[:-4] + "xyz"
        check("HMAC Signed Download Token (Rejects Tampered)", verify_download_token(tampered_token) is None)

        # 19. Malformed CSV handling
        malformed_csv = test_dir / "bad.csv"
        # Binary bytes in CSV
        malformed_csv.write_bytes(b"\x00\x80\xFF\xFE\x00\x00corrupt")
        res_bad_csv = execute_job({"operation": "csv_to_json", "input_paths": [str(malformed_csv)], "output_dir": str(out_dir)})
        check("Malformed CSV cleanly rejected", not res_bad_csv["success"] and "Malformed CSV" in res_bad_csv["error"])

        # 20. Malformed JSON handling
        malformed_json = test_dir / "bad.json"
        malformed_json.write_text("{ unquoted_key: 123,, bad }", encoding="utf-8")
        res_bad_json = execute_job({"operation": "json_to_csv", "input_paths": [str(malformed_json)], "output_dir": str(out_dir)})
        check("Malformed JSON cleanly rejected", not res_bad_json["success"] and "Malformed JSON" in res_bad_json["error"])

        # 21. Corrupted Image handling
        corrupt_img = test_dir / "corrupt.png"
        corrupt_img.write_bytes(b"\x89PNG\r\n\x1a\n\x00\x00\x00garbage_data_truncated")
        res_bad_img = execute_job({"operation": "png_to_jpg", "input_paths": [str(corrupt_img)], "output_dir": str(out_dir)})
        check("Corrupted Image cleanly rejected", not res_bad_img["success"] and "Corrupted or invalid image" in res_bad_img["error"])

        print(f"\nVerification complete: {passed}/{total} tests passed.")
        return passed == total

    finally:
        shutil.rmtree(test_dir, ignore_errors=True)

if __name__ == "__main__":
    success = run_all_tests()
    exit(0 if success else 1)
