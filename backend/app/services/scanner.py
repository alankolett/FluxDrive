import re
from typing import List, Dict, Any, Tuple
from pathlib import Path
import pymupdf as fitz
import docx

def luhn_checksum(card_number: str) -> bool:
    """Verifies credit card numbers using the standard Luhn (mod 10) algorithm."""
    digits = [int(d) for d in card_number if d.isdigit()]
    if len(digits) < 13 or len(digits) > 19:
        return False
    checksum = 0
    reverse_digits = digits[::-1]
    for i, d in enumerate(reverse_digits):
        if i % 2 == 1:
            d = d * 2
            if d > 9:
                d -= 9
        checksum += d
    return checksum % 10 == 0

def mask_value(finding_type: str, val: str) -> str:
    """Masks sensitive data for privacy-preserving UI display."""
    val = val.strip()
    if finding_type == "email":
        parts = val.split("@")
        if len(parts) == 2:
            name, domain = parts
            masked_name = name[0] + "***" if name else "***"
            return f"{masked_name}@{domain}"
        return "e***@domain"
        
    elif finding_type == "credit_card":
        digits = re.sub(r"\D", "", val)
        last4 = digits[-4:] if len(digits) >= 4 else "XXXX"
        return f"XXXX-XXXX-XXXX-{last4}"
        
    elif finding_type == "aadhaar":
        digits = re.sub(r"\D", "", val)
        last4 = digits[-4:] if len(digits) >= 4 else "XXXX"
        return f"XXXX-XXXX-{last4}"
        
    elif finding_type == "pan":
        clean = val.strip().upper()
        if len(clean) == 10:
            return f"{clean[:3]}****{clean[-1]}"
        return "XXXXX****X"
        
    elif finding_type == "phone":
        digits = re.sub(r"\D", "", val)
        last4 = digits[-4:] if len(digits) >= 4 else "XXXX"
        return f"***-***-{last4}"
        
    elif finding_type == "ipv4":
        parts = val.split(".")
        if len(parts) == 4:
            return f"{parts[0]}.{parts[1]}.*.*"
        return "*.*.*.*"
        
    elif finding_type == "url":
        if len(val) > 24:
            return val[:18] + "..."
        return val

    return "***"

def extract_text_from_file(file_path: Path, mime: str) -> str:
    """Extracts raw text content for local inspection."""
    text_chunks: List[str] = []
    
    try:
        if mime == "application/pdf":
            doc = fitz.open(file_path)
            for page in doc:
                text_chunks.append(page.get_text())
            doc.close()
            
        elif mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            doc = docx.Document(file_path)
            for p in doc.paragraphs:
                text_chunks.append(p.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text_chunks.append(cell.text)
                        
        elif mime in ["text/csv", "application/json", "text/plain"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text_chunks.append(f.read(500000))  # Up to 500KB text
    except Exception as e:
        text_chunks.append(f"Text extraction notice: {str(e)}")
        
    return "\n".join(text_chunks)

def scan_text_for_pii(text: str) -> Dict[str, Any]:
    """
    Scans extracted text for PII using regex & Luhn validation.
    Returns findings, redacted count, severity score, and masked samples.
    """
    findings: List[Dict[str, Any]] = []

    # 1. Emails
    email_matches = re.findall(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,7}\b", text)
    for m in set(email_matches):
        findings.append({
            "type": "email",
            "label": "Email Address",
            "masked": mask_value("email", m),
            "severity": "medium",
            "weight": 2,
            "verification": "Pattern verified"
        })

    # 2. Credit Cards (Luhn verified)
    potential_cards = re.findall(r"\b(?:\d[ -]*?){13,19}\b", text)
    for raw in set(potential_cards):
        clean_num = re.sub(r"\D", "", raw)
        if 13 <= len(clean_num) <= 19 and luhn_checksum(clean_num):
            findings.append({
                "type": "credit_card",
                "label": "Credit Card Number",
                "masked": mask_value("credit_card", clean_num),
                "severity": "high",
                "weight": 4,
                "verification": "Luhn algorithm verified (valid checksum)"
            })

    # 3. PAN Numbers (5 letters, 4 digits, 1 letter)
    pan_matches = re.findall(r"\b[A-Z]{5}[0-9]{4}[A-Z]\b", text, re.IGNORECASE)
    for m in set(pan_matches):
        findings.append({
            "type": "pan",
            "label": "PAN (Income Tax Identifier)",
            "masked": mask_value("pan", m),
            "severity": "high",
            "weight": 3,
            "verification": "Standard 10-character PAN structure match"
        })

    # 4. Aadhaar Numbers (12 digits, often formatted as 4 4 4)
    # Always labeled: "pattern match, not verified" as per directive
    aadhaar_matches = re.findall(r"\b[2-9]\d{3}[ -]?\d{4}[ -]?\d{4}\b", text)
    for m in set(aadhaar_matches):
        clean_num = re.sub(r"\D", "", m)
        if len(clean_num) == 12:
            findings.append({
                "type": "aadhaar",
                "label": "Aadhaar Identifier",
                "masked": mask_value("aadhaar", clean_num),
                "severity": "high",
                "weight": 4,
                "verification": "pattern match, not verified"
            })

    # 5. Phone numbers
    phone_matches = re.findall(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b", text)
    for m in set(phone_matches):
        clean = re.sub(r"\D", "", m)
        if len(clean) >= 10:
            findings.append({
                "type": "phone",
                "label": "Phone Number",
                "masked": mask_value("phone", m),
                "severity": "medium",
                "weight": 2,
                "verification": "Format pattern verified"
            })

    # 6. IPv4 Addresses
    ip_matches = re.findall(r"\b(?:(?:25[0-5]|2[0-4][0-9]|[01]?[0-9][0-9]?)\.){3}(?:25[0-5]|2[0-4][0-9]|[01]?[0-9][0-9]?)\b", text)
    for m in set(ip_matches):
        if m not in ["127.0.0.1", "0.0.0.0", "255.255.255.255"]:
            findings.append({
                "type": "ipv4",
                "label": "IPv4 Address",
                "masked": mask_value("ipv4", m),
                "severity": "medium",
                "weight": 2,
                "verification": "Octet range verified"
            })

    # 7. URLs
    url_matches = re.findall(r"https?://[^\s<>\"']+|www\.[^\s<>\"']+", text)
    for m in set(url_matches):
        findings.append({
            "type": "url",
            "label": "Web URL",
            "masked": mask_value("url", m),
            "severity": "low",
            "weight": 1,
            "verification": "URL URI scheme matched"
        })

    total_weight = sum(f["weight"] for f in findings)
    count = len(findings)
    
    if total_weight == 0:
        risk_level = "LOW"
        risk_score = 0
    elif total_weight <= 3:
        risk_level = "LOW"
        risk_score = min(total_weight * 10, 35)
    elif total_weight <= 8:
        risk_level = "MEDIUM"
        risk_score = min(35 + total_weight * 5, 69)
    else:
        risk_level = "HIGH"
        risk_score = min(70 + total_weight * 3, 100)

    return {
        "count": count,
        "risk_level": risk_level,
        "risk_score": risk_score,
        "findings": findings,
        "disclaimer": "Pattern-based detection will not catch everything; never claim perfect accuracy.",
        "text_sample_length": len(text)
    }

def sanitize_text(text: str) -> Tuple[str, int]:
    """
    Replaces identified PII with redacted placeholders before sending to AI.
    Returns: (sanitized_text, masked_field_count)
    """
    masked_count = 0
    
    def r_email(match):
        nonlocal masked_count
        masked_count += 1
        return "[REDACTED_EMAIL]"
        
    def r_card(match):
        nonlocal masked_count
        clean = re.sub(r"\D", "", match.group(0))
        if 13 <= len(clean) <= 19 and luhn_checksum(clean):
            masked_count += 1
            return "[REDACTED_CARD]"
        return match.group(0)

    def r_pan(match):
        nonlocal masked_count
        masked_count += 1
        return "[REDACTED_PAN]"

    def r_aadhaar(match):
        nonlocal masked_count
        clean = re.sub(r"\D", "", match.group(0))
        if len(clean) == 12:
            masked_count += 1
            return "[REDACTED_AADHAAR]"
        return match.group(0)

    def r_phone(match):
        nonlocal masked_count
        clean = re.sub(r"\D", "", match.group(0))
        if len(clean) >= 10:
            masked_count += 1
            return "[REDACTED_PHONE]"
        return match.group(0)

    def r_ip(match):
        nonlocal masked_count
        if match.group(0) not in ["127.0.0.1", "0.0.0.0"]:
            masked_count += 1
            return "[REDACTED_IP]"
        return match.group(0)

    cleaned = re.sub(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,7}\b", r_email, text)
    cleaned = re.sub(r"\b(?:\d[ -]*?){13,19}\b", r_card, cleaned)
    cleaned = re.sub(r"\b[A-Z]{5}[0-9]{4}[A-Z]\b", r_pan, cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\b[2-9]\d{3}[ -]?\d{4}[ -]?\d{4}\b", r_aadhaar, cleaned)
    cleaned = re.sub(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b", r_phone, cleaned)
    cleaned = re.sub(r"\b(?:(?:25[0-5]|2[0-4][0-9]|[01]?[0-9][0-9]?)\.){3}(?:25[0-5]|2[0-4][0-9]|[01]?[0-9][0-9]?)\b", r_ip, cleaned)
    
    return cleaned, masked_count
