import argparse
import hashlib
import json
import os
import sys
from pathlib import Path
from typing import Dict, Any, List
from PIL import Image
import pymupdf as fitz
import pandas as pd
import docx
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors

def compute_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()

def convert_image(in_path: Path, out_path: Path, target_fmt: str, options: dict) -> None:
    try:
        im = Image.open(in_path)
    except Exception as e:
        raise ValueError(f"Corrupted or invalid image — could not decode: {str(e)}")

    target_fmt = target_fmt.upper()
    
    if target_fmt in ["JPG", "JPEG"]:
        if im.mode in ("RGBA", "LA", "P"):
            bg = Image.new("RGB", im.size, (255, 255, 255))
            if im.mode == "P":
                im = im.convert("RGBA")
            bg.paste(im, mask=im.split()[-1] if im.mode in ("RGBA", "LA") else None)
            im = bg
        else:
            im = im.convert("RGB")
        im.save(out_path, "JPEG", quality=options.get("quality", 90), optimize=True)
        
    elif target_fmt == "PNG":
        im.save(out_path, "PNG", optimize=True)
        
    elif target_fmt == "WEBP":
        im.save(out_path, "WEBP", quality=options.get("quality", 85))
        
    elif target_fmt == "PDF":
        if im.mode != "RGB":
            bg = Image.new("RGB", im.size, (255, 255, 255))
            if "A" in im.mode:
                bg.paste(im, mask=im.split()[-1])
            else:
                bg.paste(im)
            im = bg
        im.save(out_path, "PDF", resolution=100.0)
        
    else:
        raise ValueError(f"Unsupported target image format: {target_fmt}")

def batch_images_to_pdf(input_paths: List[Path], out_path: Path) -> None:
    if not input_paths:
        raise ValueError("Batch conversion requires at least one image")
        
    opened_images = []
    try:
        for p in input_paths:
            with Image.open(p) as img:
                # Convert any format (WEBP, PNG, JPG) to clean RGB for PDF
                if img.mode in ("RGBA", "LA", "P"):
                    bg = Image.new("RGB", img.size, (255, 255, 255))
                    if img.mode == "P":
                        img = img.convert("RGBA")
                    bg.paste(img, mask=img.split()[-1] if img.mode in ("RGBA", "LA") else None)
                    opened_images.append(bg)
                else:
                    opened_images.append(img.convert("RGB"))
                    
        if not opened_images:
            raise ValueError("No valid images could be prepared for PDF compilation")
            
        first_image = opened_images[0]
        other_images = opened_images[1:]
        first_image.save(out_path, "PDF", resolution=100.0, save_all=True, append_images=other_images)
    except Exception as e:
        raise ValueError(f"Could not compile images into multi-page PDF: {str(e)}")

def pdf_to_image(in_path: Path, out_path: Path, target_fmt: str) -> None:
    try:
        doc = fitz.open(in_path)
    except Exception as e:
        raise ValueError(f"Corrupted or invalid PDF — could not parse: {str(e)}")
        
    if doc.page_count == 0:
        doc.close()
        raise ValueError("PDF document contains no pages")
        
    page = doc[0]
    pix = page.get_pixmap(dpi=150)
    
    if target_fmt.upper() == "PNG":
        pix.save(str(out_path))
    else:
        # Save as JPG
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        img.save(out_path, "JPEG", quality=90)
        
    doc.close()

def csv_to_json(in_path: Path, out_path: Path) -> None:
    try:
        df = pd.read_csv(in_path)
    except Exception as e:
        raise ValueError(f"Malformed CSV — parsing failed: {str(e)}")
        
    try:
        df.to_json(out_path, orient="records", indent=2)
    except Exception as e:
        raise ValueError(f"Could not serialize CSV data to JSON: {str(e)}")

def json_to_csv(in_path: Path, out_path: Path) -> None:
    try:
        with open(in_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except Exception as e:
        raise ValueError(f"Malformed JSON — syntax error: {str(e)}")
        
    try:
        if isinstance(data, list):
            df = pd.json_normalize(data)
        elif isinstance(data, dict):
            # Check if dict of lists or single record
            if all(isinstance(v, list) for v in data.values()):
                df = pd.DataFrame(data)
            else:
                df = pd.json_normalize([data])
        else:
            raise ValueError("JSON must contain an object or an array of objects to convert to CSV")
            
        df.to_csv(out_path, index=False)
    except Exception as e:
        raise ValueError(f"Could not structure JSON into CSV table: {str(e)}")

def docx_to_pdf(in_path: Path, out_path: Path) -> None:
    try:
        doc = docx.Document(in_path)
    except Exception as e:
        raise ValueError(f"Corrupted or invalid DOCX document: {str(e)}")
        
    pdf_doc = SimpleDocTemplate(
        str(out_path),
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    story = []
    
    # Process paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            story.append(Spacer(1, 8))
            continue
            
        style_name = p.style.name if p.style else "Normal"
        if "Heading 1" in style_name:
            story.append(Paragraph(f"<b><font size=16>{text}</font></b>", styles["Heading1"]))
            story.append(Spacer(1, 10))
        elif "Heading 2" in style_name:
            story.append(Paragraph(f"<b><font size=14>{text}</font></b>", styles["Heading2"]))
            story.append(Spacer(1, 8))
        elif "Heading 3" in style_name:
            story.append(Paragraph(f"<b><font size=12>{text}</font></b>", styles["Heading3"]))
            story.append(Spacer(1, 6))
        else:
            # Standard paragraph
            escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(escaped, styles["Normal"]))
            story.append(Spacer(1, 6))

    # Process tables
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)
        if table_data:
            t = Table(table_data)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.whitesmoke),
                ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
                ('FONTSIZE', (0,0), (-1,-1), 9),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ]))
            story.append(t)
            story.append(Spacer(1, 10))

    if not story:
        story.append(Paragraph("<i>(Empty document)</i>", styles["Italic"]))
        
    pdf_doc.build(story)

def compress_image(in_path: Path, out_path: Path, options: dict) -> None:
    try:
        im = Image.open(in_path)
    except Exception as e:
        raise ValueError(f"Corrupted or invalid image: {str(e)}")
        
    fmt = im.format if im.format else "JPEG"
    quality = options.get("quality", 60)
    
    if fmt.upper() in ["JPG", "JPEG"]:
        if im.mode != "RGB":
            im = im.convert("RGB")
        im.save(out_path, "JPEG", quality=quality, optimize=True)
    elif fmt.upper() == "PNG":
        # Compress PNG via palette quantization or optimize
        if im.mode == "RGBA":
            im.save(out_path, "PNG", optimize=True, compress_level=9)
        else:
            # Quantize for significant size savings
            try:
                quant = im.convert("RGB").quantize(colors=256)
                quant.save(out_path, "PNG", optimize=True)
            except Exception:
                im.save(out_path, "PNG", optimize=True, compress_level=9)
    elif fmt.upper() == "WEBP":
        im.save(out_path, "WEBP", quality=quality, method=6)
    else:
        im.save(out_path, format=fmt, optimize=True)

def strip_image_metadata(in_path: Path, out_path: Path) -> None:
    try:
        im = Image.open(in_path)
    except Exception as e:
        raise ValueError(f"Corrupted image — cannot read: {str(e)}")
        
    # Reconstruct pure image buffer without any metadata/EXIF dictionaries
    clean = Image.new(im.mode, im.size)
    clean.putdata(list(im.getdata()))
    
    fmt = im.format if im.format else "PNG"
    if fmt.upper() in ["JPG", "JPEG"]:
        if clean.mode != "RGB":
            clean = clean.convert("RGB")
        clean.save(out_path, "JPEG", quality=95)
    else:
        clean.save(out_path, format=fmt)

def execute_job(job_spec: Dict[str, Any]) -> Dict[str, Any]:
    operation = job_spec.get("operation")
    input_paths = [Path(p) for p in job_spec.get("input_paths", [])]
    output_dir = Path(job_spec.get("output_dir", "."))
    options = job_spec.get("options", {})
    job_id = job_spec.get("job_id", "job")
    
    if not input_paths:
        return {"success": False, "error": "No input files provided", "warnings": []}
        
    for p in input_paths:
        if not p.exists():
            return {"success": False, "error": f"Input file does not exist: {p.name}", "warnings": []}

    output_dir.mkdir(parents=True, exist_ok=True)
    first_input = input_paths[0]
    before_size = sum(p.stat().st_size for p in input_paths)
    warnings: List[str] = []
    
    # Determine output filename
    op_map = {
        "png_to_jpg": ("jpg", "image/jpeg"),
        "png_to_webp": ("webp", "image/webp"),
        "jpg_to_png": ("png", "image/png"),
        "jpg_to_webp": ("webp", "image/webp"),
        "webp_to_png": ("png", "image/png"),
        "webp_to_jpg": ("jpg", "image/jpeg"),
        "image_to_pdf": ("pdf", "application/pdf"),
        "batch_to_pdf": ("pdf", "application/pdf"),
        "pdf_to_png": ("png", "image/png"),
        "pdf_to_jpg": ("jpg", "image/jpeg"),
        "csv_to_json": ("json", "application/json"),
        "json_to_csv": ("csv", "text/csv"),
        "docx_to_pdf": ("pdf", "application/pdf"),
        "compress_image": (first_input.suffix.lstrip(".").lower(), "image/auto"),
        "strip_metadata": (first_input.suffix.lstrip(".").lower(), "image/auto"),
    }
    
    if operation not in op_map:
        return {"success": False, "error": f"Unsupported operation: {operation}", "warnings": []}
        
    out_ext, out_mime = op_map[operation]
    out_filename = f"{first_input.stem}_{operation}.{out_ext}" if operation != "batch_to_pdf" else f"combined_document.{out_ext}"
    out_path = output_dir / out_filename
    
    try:
        if operation == "png_to_jpg":
            convert_image(first_input, out_path, "JPG", options)
        elif operation == "png_to_webp":
            convert_image(first_input, out_path, "WEBP", options)
        elif operation == "jpg_to_png":
            convert_image(first_input, out_path, "PNG", options)
        elif operation == "jpg_to_webp":
            convert_image(first_input, out_path, "WEBP", options)
        elif operation == "webp_to_png":
            convert_image(first_input, out_path, "PNG", options)
        elif operation == "webp_to_jpg":
            convert_image(first_input, out_path, "JPG", options)
        elif operation == "image_to_pdf":
            convert_image(first_input, out_path, "PDF", options)
        elif operation == "batch_to_pdf":
            batch_images_to_pdf(input_paths, out_path)
        elif operation == "pdf_to_png":
            pdf_to_image(first_input, out_path, "PNG")
        elif operation == "pdf_to_jpg":
            pdf_to_image(first_input, out_path, "JPG")
        elif operation == "csv_to_json":
            csv_to_json(first_input, out_path)
        elif operation == "json_to_csv":
            json_to_csv(first_input, out_path)
        elif operation == "docx_to_pdf":
            docx_to_pdf(first_input, out_path)
        elif operation == "compress_image":
            compress_image(first_input, out_path, options)
        elif operation == "strip_metadata":
            strip_image_metadata(first_input, out_path)
            
        if not out_path.exists():
            return {"success": False, "error": "Executor completed without generating output file", "warnings": warnings}
            
        after_size = out_path.stat().st_size
        sha256 = compute_sha256(out_path)
        savings_pct = round(((before_size - after_size) / before_size) * 100, 2) if before_size > 0 else 0.0
        
        return {
            "success": True,
            "output_path": str(out_path),
            "output_filename": out_filename,
            "output_size": after_size,
            "output_mime": out_mime,
            "sha256": sha256,
            "warnings": warnings,
            "metrics": {
                "before_size": before_size,
                "after_size": after_size,
                "savings_pct": savings_pct
            }
        }
    except Exception as e:
        return {"success": False, "error": str(e), "warnings": warnings}

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="FluxDrive isolated conversion executor")
    parser.add_argument("--job", type=str, required=True, help="Job JSON specification string")
    args = parser.parse_args()
    
    try:
        spec = json.loads(args.job)
        result = execute_job(spec)
    except Exception as err:
        result = {"success": False, "error": f"Failed to parse job spec: {str(err)}", "warnings": []}
        
    print(json.dumps(result))
    sys.exit(0 if result.get("success") else 1)
